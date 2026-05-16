"""舆镜 · F4 报表导出 (Word / PowerPoint)

复刻 PDF 模板（fpdf2）的视觉与结构：
- 标题 + 副标题（生成时间 + 品牌）
- 可选首屏截图（base64 PNG）
- 章节式正文，支持基础 Markdown（# / ## / ### 标题，**粗体** 行内）

只暴露两个入口：
- ``build_docx(title, sections, images=None) -> bytes``
- ``build_pptx(title, sections, images=None) -> bytes``

调用方负责拼装 ``sections``（与 ``_build_pdf`` 完全一致的契约）。
"""

from __future__ import annotations

import base64
import io
import re
from datetime import datetime
from typing import List, Optional


# ---------------------------------------------------------------------------
# 共用工具
# ---------------------------------------------------------------------------
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_BRAND = "舆镜 YuJing"


def _decode_data_url(data_url: str) -> Optional[bytes]:
    """支持 ``data:image/png;base64,xxx`` 或裸 base64；解码失败返回 None。"""
    try:
        b64 = data_url.split(",", 1)[1] if "," in data_url else data_url
        return base64.b64decode(b64)
    except Exception:
        return None


def _now_subtitle() -> str:
    return f"{_BRAND} · {datetime.now().strftime('%Y-%m-%d %H:%M')}"


def _split_runs(line: str) -> List[tuple[str, bool]]:
    """把一行文本切成 (text, bold) 段，用于 docx run 写入。"""
    parts: List[tuple[str, bool]] = []
    cursor = 0
    for match in _BOLD_RE.finditer(line):
        if match.start() > cursor:
            parts.append((line[cursor : match.start()], False))
        parts.append((match.group(1), True))
        cursor = match.end()
    if cursor < len(line):
        parts.append((line[cursor:], False))
    return parts or [(line, False)]


# ---------------------------------------------------------------------------
# Word 导出
# ---------------------------------------------------------------------------
def build_docx(title: str, sections: list[dict], images: Optional[List[str]] = None) -> bytes:
    """生成 .docx 文档；返回二进制内容。"""
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 默认字体设为微软雅黑（中文友好）
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    # 标题
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    title_run.font.name = "微软雅黑"
    title_run.font.size = Pt(20)
    title_run.bold = True

    # 副标题
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(_now_subtitle())
    sub_run.font.size = Pt(9)
    sub_run.font.color.rgb = RGBColor(130, 130, 130)
    doc.add_paragraph()  # 空行

    # 截图
    if images:
        for data_url in images:
            raw = _decode_data_url(data_url)
            if not raw:
                continue
            try:
                img_p = doc.add_paragraph()
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.add_run().add_picture(io.BytesIO(raw), width=Inches(6.2))
            except Exception:
                continue

    # 章节
    for section in sections:
        heading = section.get("heading")
        body = section.get("body") or ""
        if heading:
            head_p = doc.add_paragraph()
            head_run = head_p.add_run(heading)
            head_run.bold = True
            head_run.font.size = Pt(13)
        for raw_line in body.splitlines():
            line = raw_line.rstrip()
            if not line:
                doc.add_paragraph()
                continue
            stripped = line.lstrip()
            # Markdown 标题降级为粗体
            if stripped.startswith("### "):
                p = doc.add_paragraph()
                run = p.add_run(stripped[4:].strip(" *"))
                run.bold = True
                run.font.size = Pt(11)
                continue
            if stripped.startswith("## "):
                p = doc.add_paragraph()
                run = p.add_run(stripped[3:].strip(" *"))
                run.bold = True
                run.font.size = Pt(13)
                continue
            if stripped.startswith("# "):
                p = doc.add_paragraph()
                run = p.add_run(stripped[2:].strip(" *"))
                run.bold = True
                run.font.size = Pt(15)
                continue
            # 普通行（行内 **粗体**）
            p = doc.add_paragraph()
            for text, is_bold in _split_runs(line):
                run = p.add_run(text)
                run.bold = is_bold

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# PPT 导出
# ---------------------------------------------------------------------------
def build_pptx(title: str, sections: list[dict], images: Optional[List[str]] = None) -> bytes:
    """生成 .pptx 文档；返回二进制内容。

    布局策略：
    - 第 1 张：封面（标题 + 副标题）
    - 若有截图：每张单独一页，居中放大
    - 每个 section 一页：章节标题 + 正文（Markdown 标题降级为粗体行）
    """
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor as _RGB

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]  # 完全空白布局
    SLIDE_W = prs.slide_width
    SLIDE_H = prs.slide_height
    MARGIN = Inches(0.6)

    def add_text_box(slide, left, top, width, height, text, *, size=14, bold=False, color=None, align_center=False):
        from pptx.enum.text import PP_ALIGN
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        if align_center:
            p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = "微软雅黑"
        run.font.size = Pt(size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = _RGB(*color)
        return tf

    # 封面
    cover = prs.slides.add_slide(blank_layout)
    add_text_box(
        cover,
        MARGIN,
        Inches(2.6),
        SLIDE_W - 2 * MARGIN,
        Inches(1.4),
        title,
        size=44,
        bold=True,
        align_center=True,
    )
    add_text_box(
        cover,
        MARGIN,
        Inches(4.2),
        SLIDE_W - 2 * MARGIN,
        Inches(0.6),
        _now_subtitle(),
        size=16,
        color=(130, 130, 130),
        align_center=True,
    )

    # 截图页
    if images:
        for data_url in images:
            raw = _decode_data_url(data_url)
            if not raw:
                continue
            slide = prs.slides.add_slide(blank_layout)
            try:
                pic_w = SLIDE_W - 2 * MARGIN
                # add_picture 需要尺寸可推断，先以宽度铺满；高度自动按比例
                slide.shapes.add_picture(
                    io.BytesIO(raw),
                    MARGIN,
                    Inches(0.6),
                    width=pic_w,
                )
            except Exception:
                continue

    # 章节页
    for section in sections:
        heading = section.get("heading") or "正文"
        body = section.get("body") or ""
        slide = prs.slides.add_slide(blank_layout)
        # 标题
        add_text_box(
            slide,
            MARGIN,
            Inches(0.4),
            SLIDE_W - 2 * MARGIN,
            Inches(0.8),
            heading,
            size=24,
            bold=True,
        )
        # 正文（一个 textbox 多 paragraph）
        body_box = slide.shapes.add_textbox(
            MARGIN,
            Inches(1.4),
            SLIDE_W - 2 * MARGIN,
            SLIDE_H - Inches(2.0),
        )
        tf = body_box.text_frame
        tf.word_wrap = True
        first = True
        for raw_line in body.splitlines():
            line = raw_line.rstrip()
            stripped = line.lstrip()
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            if not line:
                continue
            # Markdown 标题降级
            heading_level = 0
            if stripped.startswith("### "):
                heading_level, text = 3, stripped[4:].strip(" *")
            elif stripped.startswith("## "):
                heading_level, text = 2, stripped[3:].strip(" *")
            elif stripped.startswith("# "):
                heading_level, text = 1, stripped[2:].strip(" *")
            else:
                text = line

            if heading_level:
                run = p.add_run()
                run.text = text
                run.font.name = "微软雅黑"
                run.font.size = Pt({1: 18, 2: 16, 3: 14}[heading_level])
                run.font.bold = True
            else:
                for seg, is_bold in _split_runs(line):
                    run = p.add_run()
                    run.text = seg
                    run.font.name = "微软雅黑"
                    run.font.size = Pt(13)
                    run.font.bold = is_bold

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 通用辅助：HTTP 头 / 媒体类型
# ---------------------------------------------------------------------------
MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def normalize_format(fmt: Optional[str]) -> str:
    """收敛到 pdf / docx / pptx，未知值默认 pdf。"""
    f = (fmt or "pdf").strip().lower()
    return f if f in MEDIA_TYPES else "pdf"
